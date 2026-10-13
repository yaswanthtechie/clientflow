from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = "ClientFlow_Project_Starter_Guide.docx"
ACCENT = "4F6D7A"
PALE = "F4F7F8"
BORDER = "D9D9D9"
TEXT = "2F3A40"


def font(run, size=10.5, bold=False, color=TEXT, italic=False, face="Candara"):
    run.font.name = face
    run._element.rPr.rFonts.set(qn("w:ascii"), face)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), face)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic
    return run


def text(p, value, size=10.5, bold=False, color=TEXT, italic=False, face="Candara"):
    return font(p.add_run(value), size, bold, color, italic, face)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    node = tcpr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tcpr.append(node)
    node.set(qn("w:fill"), fill)
    node.set(qn("w:val"), "clear")


def border(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), BORDER)
        node.set(qn("w:space"), "0")


def pad(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    margins = tcpr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for side, value in (("top", "130"), ("start", "180"), ("bottom", "130"), ("end", "180")):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), value)
        node.set(qn("w:type"), "dxa")


def keep(paragraph):
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:keepNext"))


def no_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def header_row(row):
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(marker)


def page_number(p):
    r = p.add_run("Page ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r._r.extend([begin, instruction, end])
    font(r, 8.5, color="6B7280")


def table(doc, headers, rows, widths, size=8.9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    row = t.rows[0]
    no_split(row)
    header_row(row)
    for i, label in enumerate(headers):
        c = row.cells[i]
        c.width = Inches(widths[i])
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(c, ACCENT)
        border(c)
        pad(c)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        text(p, label, size, True, "FFFFFF")
    for ix, values in enumerate(rows):
        row = t.add_row()
        no_split(row)
        for i, value in enumerate(values):
            c = row.cells[i]
            c.width = Inches(widths[i])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade(c, "FFFFFF" if ix % 2 == 0 else PALE)
            border(c)
            pad(c)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.04
            text(p, str(value), size)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    return t


def heading(doc, label, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    keep(p)
    text(p, label, 15 if level == 1 else 11.5, True, "000000")
    return p


def body(doc, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.14
    text(p, value)
    return p


def bullet(doc, value):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    text(p, value)
    return p


def numbered(doc, value):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    text(p, value)
    return p


def code(doc, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.line_spacing = 1.03
    text(p, value, 8.65, face="Consolas")
    return p


def create():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.64)
    sec.left_margin = Inches(0.78)
    sec.right_margin = Inches(0.78)

    styles = doc.styles
    for name, face, size in (("Normal", "Candara", 10.5), ("Title", "Trebuchet MS", 27), ("Heading 1", "Candara", 15), ("Heading 2", "Candara", 11.5)):
        s = styles[name]
        s.font.name = face
        s._element.rPr.rFonts.set(qn("w:ascii"), face)
        s._element.rPr.rFonts.set(qn("w:hAnsi"), face)
        s.font.size = Pt(size)
        s.font.color.rgb = RGBColor(0, 0, 0) if name != "Normal" else RGBColor.from_string(TEXT)
        if name != "Normal":
            s.font.bold = True
    for name in ("List Bullet", "List Number"):
        styles[name].font.name = "Candara"
        styles[name].font.size = Pt(10.5)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(4)
    text(footer, "ClientFlow Starter Guide   |   ", 8.5, color="6B7280")
    page_number(footer)

    # Cover
    doc.add_paragraph().paragraph_format.space_after = Pt(42)
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(9)
    text(p, "ClientFlow Project Starter Guide", 27, True, "000000", face="Trebuchet MS")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(20)
    text(p, "From project idea to first pull request", 17, False, "000000")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(22)
    p.paragraph_format.line_spacing = 1.2
    text(p, "A practical guide for using Antigravity to plan and build ClientFlow with a React JavaScript frontend and FastAPI backend. Complete the documents and flows in this guide before asking Antigravity to generate major features.", 12)
    cover = table(doc, ["Guide covers", ""], [
        ["Planning", "PRD, TRD, application flow, user flow, backlog, and acceptance criteria."],
        ["Build workflow", "Repository setup, issue sizing, Antigravity prompts, testing, and pull requests."],
        ["Target project", "ClientFlow, a freelancer project and invoice management application."],
    ], [1.55, 4.85], 9.5)
    for c in cover.rows[0].cells:
        c.text = ""
        shade(c, "FFFFFF")
        tcpr = c._tc.get_or_add_tcPr()
        b = tcpr.first_child_found_in("w:tcBorders")
        if b is not None:
            tcpr.remove(b)
    doc.add_paragraph().paragraph_format.space_after = Pt(17)
    p = doc.add_paragraph()
    text(p, "Prepared for", 9, True, "6B7280")
    p = doc.add_paragraph()
    text(p, "Solo Developer Project Assignment", 11)
    p = doc.add_paragraph()
    text(p, "Version 1.0 | October 2026", 9, color="6B7280")
    doc.add_page_break()

    heading(doc, "1 How to Use This Guide")
    body(doc, "Use this guide in order. First set the product boundaries, then write the PRD, then make the technical plan and flows, and only then begin implementation. Antigravity is strongest when it receives clear context, a small task, an acceptance test, and a defined set of files it may change.")
    table(doc, ["Order", "Artifact", "Decision it makes"], [
        ["1", "Project boundary", "What the first release will and will not include."],
        ["2", "PRD", "Who the product is for, the problem it solves, and the required behavior."],
        ["3", "TRD", "How React, FastAPI, the database, security, and deployment will work together."],
        ["4", "Application and user flows", "What pages, actions, API calls, and outcomes exist in the user journey."],
        ["5", "Backlog", "The small, testable issues to build in a sensible order."],
        ["6", "Implementation", "A focused Antigravity request, local verification, pull request, and merge."],
    ], [0.55, 2.05, 3.8])
    heading(doc, "1.1 What to Finish Before Coding", 2)
    for item in [
        "Name the user, the problem, and the exact core workflow the first release must support.",
        "Write a short PRD and identify what is explicitly out of scope.",
        "Choose the stack once: React with JavaScript and Vite; FastAPI; PostgreSQL; SQLAlchemy; Docker.",
        "Draw the basic application and user flows in the tables provided in this guide.",
        "Create a GitHub issue for each small feature, not one issue for the whole application.",
    ]:
        bullet(doc, item)

    heading(doc, "2 Starting the Project with Antigravity")
    body(doc, "Treat Antigravity as a capable junior pair programmer. It can accelerate setup, explain code, suggest tests, and review a pull request. The developer remains responsible for choosing scope, understanding the changes, checking security, running tests, and deciding whether a change is ready to merge.")
    heading(doc, "2.1 Day Zero Setup", 2)
    table(doc, ["Step", "Action", "Done when"], [
        ["Create repository", "Create a private or public GitHub repository named clientflow.", "The default branch is main and the README describes the project."],
        ["Create workspace", "Create backend and frontend folders. Initialize FastAPI in backend and Vite React JavaScript in frontend.", "Both applications start independently."],
        ["Create docs", "Add docs/PRD.md, docs/TRD.md, docs/app-flow.md, docs/user-flow.md, and docs/backlog.md.", "Each file has a short initial version."],
        ["Create board", "Add Todo, In Progress, Review, and Done columns. Create the first 8 to 12 small issues.", "Every issue has a clear outcome and acceptance criteria."],
        ["Set environment", "Create example environment files without real secrets and add .gitignore.", "No passwords, tokens, or local database files can be committed."],
    ], [1.0, 3.15, 2.25])
    heading(doc, "2.2 Ask Antigravity in Small Units", 2)
    table(doc, ["Good request", "Why it works"], [
        ["Add the User model, registration route, and tests. Do not build login yet.", "One feature boundary, clear modules, and a testable result."],
        ["Create the React login page using the existing API client. Include loading and invalid-credential states.", "A component-level task with visible acceptance criteria."],
        ["Review this pull request for authorization bugs and missing tests. Do not edit code.", "A review request separates analysis from implementation."],
    ], [3.55, 2.85])
    heading(doc, "2.3 Prompts to Avoid", 2)
    for item in [
        "Build the whole project from scratch.",
        "Make a complete frontend and backend with all features.",
        "Fix everything that is wrong in the codebase.",
        "Make the dashboard beautiful without a defined screen, components, or acceptance criteria.",
    ]:
        bullet(doc, item)

    heading(doc, "3 Product Requirements Document PRD")
    body(doc, "The PRD describes the product decision, not the code. Keep it short enough to read before every feature. It gives Antigravity the product context it needs and prevents a technically correct feature from solving the wrong problem.")
    heading(doc, "3.1 PRD Template", 2)
    table(doc, ["Section", "What to write for ClientFlow"], [
        ["Product summary", "ClientFlow helps a freelancer organize client work, track billable time, generate invoices, and see unpaid work."],
        ["Problem", "Freelancers often track projects, hours, and invoices in separate spreadsheets or tools, making billing and follow-up difficult."],
        ["Primary user", "An independent freelancer managing several clients and projects."],
        ["Product goal", "Let a user turn completed tracked work into an accurate invoice in a single workflow."],
        ["Non goals", "Payment processing, bank integration, multi-user teams, complex tax rules, and a mobile application."],
        ["Core capabilities", "Accounts, clients, projects, tasks, time entries, invoices, dashboard, PDF and CSV export."],
        ["Success measure", "A demo user can register, add a client and project, log time, generate an invoice, and view its effect on the dashboard."],
    ], [1.45, 4.95])
    heading(doc, "3.2 PRD User Stories and Acceptance Criteria", 2)
    table(doc, ["User story", "Acceptance criteria"], [
        ["As a freelancer, I can create a client so that I can track work and billing for that client.", "Required name and email validation; client is visible only to its owner; edit and archive actions work."],
        ["As a freelancer, I can add a time entry to a task so that I can bill for completed work.", "Entry has date, duration, note, task, and billable flag; duration must be positive; totals update."],
        ["As a freelancer, I can generate an invoice from unbilled time so that I can request payment.", "Only my unbilled entries can be selected; total is calculated; entries cannot be billed twice."],
        ["As a freelancer, I can view unpaid invoices so that I know what needs follow-up.", "Dashboard shows open and overdue invoice counts and links to the invoice list."],
    ], [3.05, 3.35])
    heading(doc, "3.3 Minimum Viable Product Decision", 2)
    body(doc, "The MVP is complete when the primary user journey works end to end. A visual chart, email reminder, deployment, and extra UI polish are worthwhile only after that journey, tests, and documentation are working.")

    heading(doc, "4 Technical Requirements Document TRD")
    body(doc, "The TRD turns the PRD into an implementation contract. It identifies interfaces, data ownership, risks, and non-functional expectations. It does not need every function name; it must be clear enough that any planned feature has a natural place in the application.")
    heading(doc, "4.1 Architecture", 2)
    table(doc, ["Layer", "Responsibility", "Technology"], [
        ["Browser", "Renders pages, captures user actions, maintains route state, and displays API results.", "React, JavaScript, Vite, React Router, Bootstrap or CSS."],
        ["API client", "Sends authenticated requests and handles API errors in one place.", "Axios and an application API client module."],
        ["Backend", "Validates requests, enforces ownership, applies business rules, and serves REST endpoints.", "FastAPI, Pydantic, dependency injection, JWT, and CORS."],
        ["Data", "Stores users, clients, projects, tasks, time entries, invoices, and migrations.", "PostgreSQL, SQLAlchemy, Alembic."],
        ["Runtime", "Runs frontend, backend, and database consistently in development.", "Docker and Docker Compose."],
    ], [1.0, 3.75, 1.65])
    heading(doc, "4.2 System Flow", 2)
    table(doc, ["Step", "Request and response flow"], [
        ["1", "User selects a page or submits a React form."],
        ["2", "React calls the FastAPI endpoint through the shared Axios client."],
        ["3", "FastAPI checks the JWT, validates input with Pydantic, and confirms resource ownership."],
        ["4", "A service applies business rules and SQLAlchemy reads or writes PostgreSQL."],
        ["5", "FastAPI returns JSON or a download response; React updates the screen state."],
    ], [0.55, 5.85])
    heading(doc, "4.3 Data and API Rules", 2)
    for item in [
        "Every user-owned record must have a user_id or be reachable only through a user-owned parent record.",
        "A route must never accept a user_id from the browser as proof of ownership; obtain the current user from the JWT dependency.",
        "React should not calculate invoice totals as the source of truth. The backend service calculates and returns totals.",
        "Use one Axios API client to attach the access token and translate expected API failures into user-facing messages.",
        "Enable CORS only for the local React development origin and the final deployed frontend origin.",
    ]:
        bullet(doc, item)
    heading(doc, "4.4 Suggested API Contract", 2)
    table(doc, ["Area", "Representative endpoints"], [
        ["Authentication", "POST /auth/register, POST /auth/login, GET /auth/me"],
        ["Clients", "GET /clients, POST /clients, GET /clients/{id}, PATCH /clients/{id}"],
        ["Projects and tasks", "GET or POST /clients/{id}/projects, GET or POST /projects/{id}/tasks"],
        ["Time entries", "GET /tasks/{id}/time-entries, POST /tasks/{id}/time-entries"],
        ["Invoices", "POST /invoices/from-time-entries, GET /invoices, PATCH /invoices/{id}/status"],
        ["Reporting", "GET /dashboard/summary, GET /reports/monthly-income, GET /export/time-entries.csv"],
    ], [1.45, 4.95])

    heading(doc, "5 Application Flow")
    body(doc, "Application flow maps the screens and the system behavior behind them. It is useful for choosing React routes, components, API endpoints, and empty/error states before UI implementation begins.")
    table(doc, ["Screen and route", "User action", "Backend action", "Next state"], [
        ["Login /login", "Submit email and password.", "POST /auth/login returns JWT and user data.", "Store session safely for the app and navigate to dashboard."],
        ["Dashboard /dashboard", "Open dashboard.", "GET /dashboard/summary.", "Show income, active projects, unpaid and overdue invoices."],
        ["Clients /clients", "Create or choose a client.", "GET or POST /clients.", "Refresh list or open client details."],
        ["Project /projects/:id", "Create task or log time.", "GET project, tasks, and time entries; POST new item.", "Update task list and project totals."],
        ["Invoice /invoices/new", "Select unbilled entries and generate invoice.", "POST /invoices/from-time-entries.", "Open invoice details with calculated totals."],
        ["Invoice /invoices/:id", "Change status or download PDF.", "PATCH status or GET PDF response.", "Show updated status or browser download."],
    ], [1.25, 1.65, 2.2, 1.3], 8.25)
    heading(doc, "5.1 React Page Inventory", 2)
    table(doc, ["Page", "Core components", "States to design"], [
        ["Authentication", "LoginForm, RegisterForm", "Initial, submitting, validation error, invalid credentials."],
        ["Dashboard", "SummaryCards, InvoiceList, ProjectList", "Loading, no data, API error, populated."],
        ["Clients", "ClientForm, ClientTable, ClientDetail", "Create, edit, archive confirmation, empty list."],
        ["Projects", "ProjectForm, TaskList, TaskForm, TimeEntryForm", "Active, completed, no tasks, invalid input."],
        ["Invoices", "InvoiceBuilder, InvoiceDetail, InvoiceStatusBadge", "Draft, sent, paid, overdue, no billable time."],
    ], [1.15, 3.0, 2.25])

    heading(doc, "6 User Flow")
    body(doc, "User flow focuses on the human journey. Each step should be easy to explain in a demo and should have a clear failure path. The primary flow below defines the MVP.")
    table(doc, ["Step", "User goal", "Success path", "Failure or edge case"], [
        ["1", "Create account", "Register, validate form, then log in.", "Duplicate email or invalid password shows a useful message."],
        ["2", "Add client", "Save a client with a default hourly rate.", "Required fields are missing or client is archived."],
        ["3", "Plan work", "Create project and tasks with estimates and due dates.", "Invalid dates or task state transition is rejected."],
        ["4", "Track work", "Add a positive-duration time entry against a task.", "Task not found, invalid duration, or entry already billed."],
        ["5", "Bill work", "Select unbilled entries and generate a draft invoice.", "No entries selected or an entry is already billed."],
        ["6", "Track payment", "Mark invoice sent then paid; dashboard updates.", "Illegal state change is rejected and current status remains visible."],
    ], [0.45, 1.15, 2.65, 2.15], 8.35)
    heading(doc, "6.1 First Demo Script", 2)
    for item in [
        "Log in using seeded demo credentials or register a new user.",
        "Create a client and open its project.",
        "Create a task and record a time entry.",
        "Generate a draft invoice from the unbilled entry.",
        "Download the invoice PDF and change its status to sent or paid.",
        "Return to the dashboard and explain the updated figures.",
    ]:
        numbered(doc, item)

    heading(doc, "7 Backlog and Issue Design")
    body(doc, "A backlog is a build order, not just a feature list. Keep each issue small enough to complete, test, review, and merge in one focused session or day. Write acceptance criteria before starting implementation.")
    table(doc, ["Epic", "First issues", "Dependency"], [
        ["Foundation", "Repository, Docker Compose, FastAPI base, Vite React base, environment template.", "None."],
        ["Authentication", "User model, register endpoint, login endpoint, React login page, route protection.", "Foundation."],
        ["Client and project", "Client CRUD, project CRUD, client and project pages, ownership tests.", "Authentication."],
        ["Task and time", "Task CRUD, time entry CRUD, project totals, React task and time forms.", "Client and project."],
        ["Invoices", "Invoice data model, invoice generation service, list and detail pages, PDF export.", "Task and time."],
        ["Reporting", "Dashboard API, React summary cards, export, demo data, quality pass.", "Invoices."],
    ], [1.15, 3.85, 1.35], 8.35)
    heading(doc, "7.1 GitHub Issue Template", 2)
    code(doc, "Title: Add client creation endpoint\n\nGoal: An authenticated user can create a client they own.\n\nAcceptance criteria:\n- POST /clients accepts name, email, phone, and hourly_rate.\n- The created record contains the current authenticated user ID.\n- Invalid email and negative hourly_rate return validation errors.\n- Tests cover successful creation and unauthenticated access.\n- API documentation is clear in /docs.\n\nOut of scope: Client page UI and client editing.")

    heading(doc, "8 Pull Request Workflow")
    body(doc, "Open a pull request for every completed issue, even when working alone. The PR is a checkpoint: it records why the change exists, isolates the diff, and gives Antigravity an exact unit to review.")
    table(doc, ["Stage", "Developer action", "Antigravity request"], [
        ["Plan", "Create an issue and a branch named feature/client-create.", "Read the issue and identify files, risks, and tests before coding."],
        ["Build", "Implement one scoped feature and run relevant tests.", "Implement only the issue requirements; explain each changed file."],
        ["Review", "Open a PR with a concise summary and screenshots for UI work.", "Review for broken authorization, missing errors, duplication, and missing tests. Do not edit."],
        ["Merge", "Fix findings, rerun tests, then merge to main.", "Summarize what is verified and which risks remain."],
    ], [0.8, 2.75, 2.8])
    heading(doc, "8.1 Pull Request Description Template", 2)
    code(doc, "What changed\n- Add client creation API route and validation.\n- Add tests for success and unauthenticated requests.\n\nHow to test\n- Run pytest.\n- Use POST /clients in Swagger with a valid JWT.\n\nAcceptance criteria\n- [x] Current user owns created client\n- [x] Validation errors are returned\n- [x] Tests pass\n\nOut of scope\n- React client form")

    heading(doc, "9 Prompt Library for Antigravity")
    body(doc, "Copy these patterns and replace only the feature-specific details. Always attach or summarize the relevant PRD, TRD, API contract, and acceptance criteria. Ask for an explanation after code generation, then test the output yourself.")
    heading(doc, "9.1 Planning Prompt", 2)
    code(doc, "We are building ClientFlow with React JavaScript and FastAPI. Review this issue: [paste issue]. Using the PRD and TRD, propose a small implementation plan. List files to create or modify, API contract changes, test cases, and risks. Do not write code yet.")
    heading(doc, "9.2 Implementation Prompt", 2)
    code(doc, "Implement this one issue in the existing ClientFlow repository: [paste issue]. Respect the current React and FastAPI folder structure. Do not modify unrelated files. Add or update tests. Before coding, list the files you will touch. After coding, explain every changed file, provide test commands, and list anything not implemented.")
    heading(doc, "9.3 Review Prompt", 2)
    code(doc, "Review the current branch against this acceptance criteria: [paste criteria]. Focus on authorization, validation, API and React error states, data integrity, and missing tests. Do not edit files. Report findings by severity and include an exact explanation for each one.")

    heading(doc, "10 Daily Operating Routine")
    table(doc, ["Time", "Routine"], [
        ["Start of day", "Read the current GitHub issue, PRD section, TRD note, and acceptance criteria. Ask Antigravity for a plan only if the task is unclear."],
        ["During build", "Work on one branch and one issue. Run the feature locally after each meaningful change."],
        ["Before PR", "Read the diff, verify no secrets exist, test success and failure cases, and update documentation if needed."],
        ["After merge", "Run the main branch, update the board, and write down a short lesson or blocker for the next day."],
    ], [1.15, 5.25])
    heading(doc, "10.1 Definition of Done", 2)
    for item in [
        "The user story and acceptance criteria are met.",
        "React handles loading, empty, success, and useful failure states where applicable.",
        "FastAPI validates input and enforces ownership on every protected resource.",
        "Relevant tests pass and the feature was checked manually in the browser or Swagger.",
        "The pull request explains the change and contains no unrelated code or secrets.",
        "The developer can explain the data flow from React action to database result and back to the screen.",
    ]:
        bullet(doc, item)

    heading(doc, "11 First Week Checklist")
    body(doc, "The first week should create a dependable base, not a long list of half-built features. Finish this checklist before starting client and project modules.")
    table(doc, ["Checklist", "Target"], [
        ["Repository and documentation", "README and initial PRD, TRD, application flow, user flow, and backlog exist."],
        ["Backend", "FastAPI starts, PostgreSQL connects, migration system works, and health endpoint responds."],
        ["Frontend", "Vite React JavaScript application starts and can call the backend health endpoint."],
        ["Authentication", "Registration, login, JWT, protected API route, route guard, and login screen work."],
        ["Quality", "Tests cover authentication basics; Docker Compose starts the application from documented steps."],
    ], [2.1, 4.3])
    body(doc, "When this checklist is complete, the project is ready for feature development. Keep the PRD and TRD as living documents: update them when a deliberate product or technical decision changes, not after every small code edit.")

    doc.save(OUT)


if __name__ == "__main__":
    create()
